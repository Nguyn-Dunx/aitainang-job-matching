"""Tầng 1 cho D2 — đọc CV từ PDF/DOCX thành text thô.

Chiến lược:
- PDF: PyMuPDF (fitz) với sort=True để sắp xếp block theo thứ tự đọc (quan trọng
  với CV nhiều cột). Nếu chất lượng text quá thấp (CV scan/ảnh) thì fallback Docling
  (optional dependency) hoặc cảnh báo cần OCR.
- DOCX: python-docx (đọc cả bảng — CV hay đặt nội dung trong table cell).
- KHÔNG dùng cho JD: JD đã có cấu trúc sẵn từ data/processed/jds.json.

ParseResult.warnings luôn được trả về để tầng trên quyết định human-in-the-loop.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path

log = logging.getLogger(__name__)

MIN_CHARS_PER_PAGE = 100  # dưới ngưỡng này: nghi CV dạng ảnh scan


@dataclass
class ParseResult:
    text: str
    page_count: int
    method: str  # "pymupdf" | "docling" | "python-docx"
    warnings: list[str] = field(default_factory=list)


def _parse_pdf_pymupdf(path: Path) -> ParseResult:
    import fitz  # PyMuPDF

    doc = fitz.open(path)
    pages = [page.get_text("text", sort=True) for page in doc]
    result = ParseResult(
        text="\n\n".join(pages).strip(),
        page_count=len(doc),
        method="pymupdf",
    )
    doc.close()
    return result


def _parse_pdf_docling(path: Path) -> ParseResult:
    """Fallback cho layout phức tạp. Docling là optional dependency."""
    try:
        from docling.document_converter import DocumentConverter
    except ImportError as e:
        raise RuntimeError("Docling chưa được cài (pip install docling)") from e
    converted = DocumentConverter().convert(str(path))
    text = converted.document.export_to_markdown()
    return ParseResult(text=text.strip(), page_count=0, method="docling")


def _parse_docx(path: Path) -> ParseResult:
    import docx

    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    # CV thường đặt nội dung trong bảng — đọc cả table cell
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs if p.text.strip())
    return ParseResult(text="\n".join(parts), page_count=1, method="python-docx")


def parse_cv(path: str | Path, use_docling_fallback: bool = True) -> ParseResult:
    """Điểm vào duy nhất của Tầng 1. Nhận .pdf hoặc .docx."""
    path = Path(path)
    suffix = path.suffix.lower()
    if not path.exists():
        raise FileNotFoundError(path)

    if suffix == ".docx":
        return _parse_docx(path)
    if suffix != ".pdf":
        raise ValueError(f"Định dạng chưa hỗ trợ: {suffix} (chỉ .pdf/.docx)")

    result = _parse_pdf_pymupdf(path)
    chars_per_page = len(result.text) / max(result.page_count, 1)
    if chars_per_page < MIN_CHARS_PER_PAGE:
        result.warnings.append(
            f"Text quá ít ({chars_per_page:.0f} ký tự/trang) — nghi CV dạng ảnh scan, cần OCR."
        )
        if use_docling_fallback:
            try:
                fallback = _parse_pdf_docling(path)
                if len(fallback.text) > len(result.text):
                    fallback.warnings.append("Đã dùng Docling fallback cho layout phức tạp.")
                    return fallback
            except RuntimeError as e:
                result.warnings.append(str(e))
    return result
