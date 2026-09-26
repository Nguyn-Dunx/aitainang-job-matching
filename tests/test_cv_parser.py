"""Test Tầng 1: parse CV PDF/DOCX. Fixture = 3 CV synthetic (scripts/make_synthetic_cvs.py)."""

from pathlib import Path

import pytest

from app.services.cv_parser import parse_cv

FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture(scope="module", autouse=True)
def ensure_fixtures():
    if not (FIXTURES / "cv_single_column.pdf").exists():
        pytest.skip("Chưa có fixture — chạy: python scripts/make_synthetic_cvs.py")


def test_parse_single_column_day_du_section():
    r = parse_cv(FIXTURES / "cv_single_column.pdf")
    assert r.method == "pymupdf"
    assert r.page_count == 1
    assert not r.warnings
    text = r.text
    for keyword in ["NGUYEN VAN AN", "KY NANG", "HOC VAN", "KINH NGHIEM", "DU AN", "Python", "Django"]:
        assert keyword in text, f"Thiếu: {keyword}"


def test_parse_two_column_giu_du_noi_dung_hai_cot():
    r = parse_cv(FIXTURES / "cv_two_column.pdf")
    text = r.text
    # Nội dung cả 2 cột đều phải xuất hiện (thứ tự tương đối do sort=True quyết định)
    for keyword in ["TRAN THI BICH", "LIEN HE", "KY NANG", "KINH NGHIEM", "Spring Boot", "Kafka"]:
        assert keyword in text, f"Thiếu: {keyword}"


def test_parse_cv_thieu_muc_khong_crash():
    r = parse_cv(FIXTURES / "cv_missing_sections.pdf")
    assert "LE MINH CUONG" in r.text
    assert "Pandas" in r.text


def test_dinh_dang_khong_ho_tro_bao_loi(tmp_path):
    fake = tmp_path / "cv.txt"
    fake.write_text("hello", encoding="utf-8")
    with pytest.raises(ValueError):
        parse_cv(fake)


def test_file_khong_ton_tai_bao_loi():
    with pytest.raises(FileNotFoundError):
        parse_cv(FIXTURES / "khong_co.pdf")


def test_docx(tmp_path):
    import docx

    path = tmp_path / "cv.docx"
    d = docx.Document()
    d.add_heading("PHAM VAN DUC", level=1)
    d.add_paragraph("Kỹ năng: Java, Spring Boot, MySQL")
    d.save(path)
    r = parse_cv(path)
    assert r.method == "python-docx"
    assert "PHAM VAN DUC" in r.text
    assert "Spring Boot" in r.text
